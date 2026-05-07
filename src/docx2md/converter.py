"""Core conversion logic."""

import json
from pathlib import Path
from typing import Optional

import markdown

try:
    import markitdown
except ImportError:
    markitdown = None

try:
    import weasyprint
except ImportError:
    weasyprint = None

try:
    from docx import Document
except ImportError:
    Document = None

try:
    from ebooklib import epub
except ImportError:
    epub = None


class Converter:
    """Handles document conversion via Markdown intermediate."""

    def __init__(self):
        self._markdown_converter = None

    def parse_to_markdown(self, input_path: str) -> str:
        """Parse input file to Markdown using markitdown."""
        if markitdown is None:
            raise RuntimeError("markitdown not installed")

        result = markitdown.MarkItDown().convert(input_path)
        return result.text_content

    def convert_markdown_to(self, markdown_content: str, output_format: str, output_path: str) -> None:
        """Convert Markdown to target format."""
        method = f"_to_{output_format.lower()}"
        if not hasattr(self, method):
            raise ValueError(f"Unsupported output format: {output_format}")
        getattr(self, method)(markdown_content, output_path)

    def _to_html(self, md: str, output_path: str) -> None:
        """Convert Markdown to HTML."""
        html = markdown.markdown(md)
        wrapper = f"""<!DOCTYPE html>
<html>
<head><meta charset="utf-8"><title></title></head>
<body>{html}</body>
</html>"""
        Path(output_path).write_text(wrapper, encoding="utf-8")

    def _to_txt(self, md: str, output_path: str) -> None:
        """Convert Markdown to plain text."""
        import re
        text = re.sub(r'[#*_`\[\]()]', '', md)
        text = re.sub(r'\n{3,}', '\n\n', text)
        Path(output_path).write_text(text, encoding="utf-8")

    def _to_json(self, md: str, output_path: str) -> None:
        """Convert Markdown to JSON structure."""
        lines = md.split('\n')
        result = {"content": []}
        for line in lines:
            if line.startswith('# '):
                result["content"].append({"type": "heading1", "text": line[2:]})
            elif line.startswith('## '):
                result["content"].append({"type": "heading2", "text": line[3:]})
            elif line.startswith('### '):
                result["content"].append({"type": "heading3", "text": line[4:]})
            elif line.strip():
                result["content"].append({"type": "paragraph", "text": line})
        Path(output_path).write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")

    def _to_pdf(self, md: str, output_path: str) -> None:
        """Convert Markdown to PDF via HTML."""
        if weasyprint is None:
            raise RuntimeError("weasyprint not installed")
        html = markdown.markdown(md)
        wrapper = f"""<!DOCTYPE html>
<html>
<head><meta charset="utf-8"><title></title></head>
<body>{html}</body>
</html>"""
        weasyprint.HTML(string=wrapper).write_pdf(output_path)

    def _to_docx(self, md: str, output_path: str) -> None:
        """Convert Markdown to DOCX."""
        if Document is None:
            raise RuntimeError("python-docx not installed")
        doc = Document()
        for line in md.split('\n'):
            if line.startswith('# '):
                doc.add_heading(line[2:], level=1)
            elif line.startswith('## '):
                doc.add_heading(line[3:], level=2)
            elif line.startswith('### '):
                doc.add_heading(line[4:], level=3)
            elif line.strip():
                doc.add_paragraph(line)
        doc.save(output_path)

    def _to_epub(self, md: str, output_path: str) -> None:
        """Convert Markdown to EPUB."""
        if epub is None:
            raise RuntimeError("ebooklib not installed")
        book = epub.EpubBook()
        book.set_identifier('docx2md')
        book.set_title('Converted Document')
        book.set_language('en')

        chapters = []
        chapter = epub.EpubHtml(title='Main', file_name='chapter.xhtml', lang='en')
        chapter.content = f'<html><body>{markdown.markdown(md)}</body></html>'
        book.add_item(chapter)
        chapters.append(chapter)

        book.toc = tuple(chapters)
        book.spine = ['nav'] + chapters

        epub.write_epub(output_path, book)

    def _to_markdown(self, md: str, output_path: str) -> None:
        """Save as Markdown."""
        Path(output_path).write_text(md, encoding="utf-8")


def convert(input_path: str, output_path: str) -> None:
    """Single-file conversion."""
    converter = Converter()
    md = converter.parse_to_markdown(input_path)
    output_format = Path(output_path).suffix.lower().lstrip('.')
    if output_format == 'md':
        output_format = 'markdown'
    converter.convert_markdown_to(md, output_format, output_path)